"""Writes the two-sheet `.xlsx` workbook for `export_excel` (ticket 10, D14):
`Data` (formatted rows, sharing column typing with `create_table`) and `Details`
(provenance as key/value rows, mirroring `OutputProvenance`). Sheet names are
fixed, never derived from the report title."""

from __future__ import annotations

import io
import re
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.schemas.chats import OutputProvenance, TableColumn
from app.utils.column_typing import excel_number_format

XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_SLUG_COLLAPSE_RE = re.compile(r"[^a-z0-9]+")


def slugify(text: str) -> str:
    """`"Top vendors!"` -> `"top-vendors"`; empty/punctuation-only falls back to
    `"report"` so the filename formula never degenerates to a bare `-{date}.xlsx`."""
    slug = _SLUG_COLLAPSE_RE.sub("-", text.strip().lower()).strip("-")
    return slug or "report"


def default_export_filename(store_id: Optional[str], title: str, generated_at: datetime) -> str:
    """`{store_id}-{slug(title)}-{YYYY-MM-DD}.xlsx` (D14) — no store-name lookup."""
    return f"{store_id or 'store'}-{slugify(title)}-{generated_at:%Y-%m-%d}.xlsx"


def sanitize_export_filename(filename: str) -> str:
    """Takes the basename (dropping any path/traversal segments) and forces a
    `.xlsx` extension, discarding whatever extension was supplied."""
    name = filename.strip().replace("\\", "/")
    basename = name.rsplit("/", 1)[-1].lstrip(".")
    stem = Path(basename).stem.strip() if basename else ""
    return f"{stem or 'export'}.xlsx"


def _excel_safe(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime) and value.tzinfo is not None:
        # openpyxl rejects tz-aware datetimes outright; Excel has no timezone concept.
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def _display_width(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, (datetime, date)):
        return len(value.isoformat())
    return len(str(value))


def build_report_workbook(
    columns: List[TableColumn],
    rows: List[Dict[str, Any]],
    provenance: OutputProvenance,
) -> bytes:
    """Full (untruncated) `rows` in, workbook bytes out. `rows` are raw stored-result
    dicts (native Python types); Excel needs no separate JSON-safety pass."""
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Data"

    header_font = Font(bold=True)
    for col_idx, column in enumerate(columns, start=1):
        cell = data_sheet.cell(row=1, column=col_idx, value=column.label)
        cell.font = header_font

    number_formats = [
        excel_number_format(column.key, column.type, (r.get(column.key) for r in rows))
        for column in columns
    ]

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, column in enumerate(columns, start=1):
            cell = data_sheet.cell(
                row=row_idx, column=col_idx, value=_excel_safe(row.get(column.key))
            )
            fmt = number_formats[col_idx - 1]
            if fmt:
                cell.number_format = fmt

    data_sheet.freeze_panes = "A2"
    last_col_letter = get_column_letter(max(len(columns), 1))
    last_row = len(rows) + 1
    data_sheet.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    for col_idx, column in enumerate(columns, start=1):
        widest = max(
            [len(column.label)] + [_display_width(r.get(column.key)) for r in rows]
        )
        data_sheet.column_dimensions[get_column_letter(col_idx)].width = min(
            max(widest + 2, 10), 60
        )

    details_sheet = workbook.create_sheet("Details")
    details_sheet.append(["sql", provenance.sql])
    for step in provenance.transforms:
        details_sheet.append([f"transform: {step.transform}", str(step.params)])
    details_sheet.append(["row_count", provenance.row_count])
    details_sheet.append(["generated_at", provenance.generated_at.isoformat()])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
