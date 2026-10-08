"""Pure-function tests for the export_excel workbook writer (ticket 10, D14):
sheet names, header bold + freeze + autofilter, per-type cell formats, `_pct`
columns as plain numbers, filename formula, and filename sanitisation. Workbook
assertions read the written bytes back with `openpyxl` (testing-decisions seam)."""

import io
from datetime import datetime, timezone

import openpyxl
import pytest

from app.schemas.chats import OutputProvenance, TableColumn, TransformStep
from app.utils.xlsx_export import (
    build_report_workbook,
    default_export_filename,
    sanitize_export_filename,
    slugify,
)


def _columns():
    return [
        TableColumn(key="vendor_name", label="Vendor Name", type="string"),
        TableColumn(key="order_total", label="Order Total", type="currency"),
        TableColumn(key="order_date", label="Order Date", type="date"),
        TableColumn(key="quantity", label="Quantity", type="number"),
        TableColumn(key="share_pct", label="Share Pct", type="number"),
    ]


def _rows():
    return [
        {
            "vendor_name": "Acme",
            "order_total": 100.5,
            "order_date": datetime(2026, 1, 15, tzinfo=timezone.utc),
            "quantity": 3,
            "share_pct": 23.5,
        },
        {
            "vendor_name": "Globex",
            "order_total": 200.0,
            "order_date": datetime(2026, 2, 1, tzinfo=timezone.utc),
            "quantity": 5,
            "share_pct": 76.5,
        },
    ]


def _provenance(**overrides):
    defaults = dict(
        sql="SELECT * FROM vendors",
        transforms=[TransformStep(transform="top_n", params={"n": 10})],
        row_count=2,
        generated_at=datetime(2026, 9, 16, tzinfo=timezone.utc),
    )
    defaults.update(overrides)
    return OutputProvenance(**defaults)


def _load(workbook_bytes: bytes):
    return openpyxl.load_workbook(io.BytesIO(workbook_bytes))


def test_workbook_has_exactly_data_and_details_sheets():
    wb_bytes = build_report_workbook(_columns(), _rows(), _provenance())
    wb = _load(wb_bytes)
    assert wb.sheetnames == ["Data", "Details"]


def test_data_sheet_header_is_bold_and_humanised():
    wb_bytes = build_report_workbook(_columns(), _rows(), _provenance())
    wb = _load(wb_bytes)
    data = wb["Data"]
    header_values = [cell.value for cell in data[1]]
    assert header_values == ["Vendor Name", "Order Total", "Order Date", "Quantity", "Share Pct"]
    assert all(cell.font.bold for cell in data[1])


def test_data_sheet_freezes_header_row_and_sets_autofilter():
    wb_bytes = build_report_workbook(_columns(), _rows(), _provenance())
    wb = _load(wb_bytes)
    data = wb["Data"]
    assert data.freeze_panes == "A2"
    assert data.auto_filter.ref == "A1:E3"


def test_data_sheet_writes_all_rows():
    wb_bytes = build_report_workbook(_columns(), _rows(), _provenance())
    wb = _load(wb_bytes)
    data = wb["Data"]
    assert [cell.value for cell in data[2]][:2] == ["Acme", 100.5]
    assert [cell.value for cell in data[3]][:2] == ["Globex", 200.0]


def test_cell_number_formats_per_inferred_type():
    wb_bytes = build_report_workbook(_columns(), _rows(), _provenance())
    wb = _load(wb_bytes)
    data = wb["Data"]
    # Row 2: vendor_name, order_total, order_date, quantity, share_pct
    assert data.cell(row=2, column=2).number_format == "$#,##0.00"
    assert data.cell(row=2, column=3).number_format == "yyyy-mm-dd"
    assert data.cell(row=2, column=4).number_format == "#,##0"


def test_share_pct_column_gets_plain_number_format_not_percent():
    wb_bytes = build_report_workbook(_columns(), _rows(), _provenance())
    wb = _load(wb_bytes)
    data = wb["Data"]
    fmt = data.cell(row=2, column=5).number_format
    assert fmt == "#,##0.0"
    assert "%" not in fmt


def test_details_sheet_mirrors_provenance_with_one_row_per_transform_step():
    wb_bytes = build_report_workbook(_columns(), _rows(), _provenance())
    wb = _load(wb_bytes)
    details = wb["Details"]
    rows = [[cell.value for cell in row] for row in details.iter_rows()]
    assert rows[0] == ["sql", "SELECT * FROM vendors"]
    assert rows[1][0] == "transform: top_n"
    assert "10" in rows[1][1]
    assert rows[2] == ["row_count", 2]
    assert rows[3][0] == "generated_at"


def test_details_sheet_has_one_row_per_chained_transform():
    provenance = _provenance(
        transforms=[
            TransformStep(transform="pivot", params={}),
            TransformStep(transform="top_n", params={"n": 5}),
        ]
    )
    wb_bytes = build_report_workbook(_columns(), _rows(), provenance)
    wb = _load(wb_bytes)
    details = wb["Details"]
    transform_rows = [
        row[0].value for row in details.iter_rows() if str(row[0].value).startswith("transform:")
    ]
    assert transform_rows == ["transform: pivot", "transform: top_n"]


def test_empty_result_set_still_writes_a_valid_workbook():
    wb_bytes = build_report_workbook(_columns(), [], _provenance(row_count=0))
    wb = _load(wb_bytes)
    assert wb.sheetnames == ["Data", "Details"]
    assert wb["Data"].auto_filter.ref == "A1:E1"


@pytest.mark.parametrize(
    "title,expected_slug",
    [
        ("Top vendors", "top-vendors"),
        ("Sales by State!!", "sales-by-state"),
        ("  ", "report"),
    ],
)
def test_slugify(title, expected_slug):
    assert slugify(title) == expected_slug


def test_default_export_filename_formula():
    generated_at = datetime(2026, 9, 16, tzinfo=timezone.utc)
    assert (
        default_export_filename("store-9", "Top vendors", generated_at)
        == "store-9-top-vendors-2026-09-16.xlsx"
    )


def test_sanitize_export_filename_strips_traversal_and_forces_xlsx():
    assert sanitize_export_filename("../evil.txt") == "evil.xlsx"


def test_sanitize_export_filename_replaces_existing_extension():
    assert sanitize_export_filename("my report.csv") == "my report.xlsx"


def test_sanitize_export_filename_handles_blank_name():
    assert sanitize_export_filename("   ") == "export.xlsx"
